"""Test build_workbook.py builds all sheets from SYNTHETIC fixtures."""
import os
import sys
import tempfile

import pandas as pd
import yaml
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

from openpyxl import Workbook  # noqa: E402

import ingest          # noqa: E402
import analyze         # noqa: E402
import build_workbook  # noqa: E402

FIX = os.path.join(ROOT, "tests", "fixtures")
CFG = ingest.load_config()


def _fake_template(path):
    """Mimic the official template: title/notes rows, header at row 4, a sample row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "VOC Collection"
    ws["A1"] = "Paytm UPI VOC collection template"
    ws["A2"] = "* Feel free to add more columns"
    hdr = ["VOC ID", "Date", "Respondent type", "City / area", "Profile / category",
           "UPI apps used", "Primary UPI app used", "Why was it chosen? (exact words)",
           "When or why is Paytm used?", "Need, barrier or motivation",
           "Opportunity / idea", "Key quote"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=j, value=h)
    ws.cell(row=5, column=1, value="VOC-001")   # sample row to be cleared
    ws.cell(row=5, column=3, value="Consumer")
    wb.save(path)


def test_workbook_sheets_and_rows():
    df = pd.read_csv(os.path.join(FIX, "form_responses.SYNTHETIC.csv"))
    fr, fp = ingest.build_from_form(df, 500)
    ir, ip = ingest.build_from_interviews(os.path.join(FIX, "interviews"), 500)
    respondents = pd.concat([fr, ir], ignore_index=True)
    payments = pd.concat([fp, ip], ignore_index=True)
    metrics = analyze.compute_metrics(respondents, payments, CFG)
    codebook = yaml.safe_load(open(os.path.join(ROOT, "docs", "codebook.yaml"), encoding="utf-8"))

    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "T_VOC.xlsx")
        build_workbook.build_workbook(
            respondents, payments, metrics, codebook,
            template_path=os.path.join(d, "missing_template.xlsx"),
            out_path=out, charts_dir=os.path.join(d, "charts"))
        assert os.path.exists(out)
        wb = load_workbook(out)
        for sheet in ["VOC Collection", "Payments", "Summary", "Codebook", "Method"]:
            assert sheet in wb.sheetnames
        # VOC Collection has header + one row per respondent
        voc = wb["VOC Collection"]
        assert voc.max_row == 1 + len(respondents)
        # Payments has header + one row per payment
        pay = wb["Payments"]
        assert pay.max_row == 1 + len(payments)


def test_template_mode_preserves_and_appends():
    df = pd.read_csv(os.path.join(FIX, "form_responses.SYNTHETIC.csv"))
    fr, fp = ingest.build_from_form(df, 500)
    ir, ip = ingest.build_from_interviews(os.path.join(FIX, "interviews"), 500)
    respondents = pd.concat([fr, ir], ignore_index=True)
    payments = pd.concat([fp, ip], ignore_index=True)
    metrics = analyze.compute_metrics(respondents, payments, CFG)
    codebook = yaml.safe_load(open(os.path.join(ROOT, "docs", "codebook.yaml"), encoding="utf-8"))

    with tempfile.TemporaryDirectory() as d:
        template = os.path.join(d, "template.xlsx")
        _fake_template(template)
        out = os.path.join(d, "Aquaholics_VOC.xlsx")
        build_workbook.build_workbook(
            respondents, payments, metrics, codebook,
            template_path=template, out_path=out,
            charts_dir=os.path.join(d, "charts"))
        wb = load_workbook(out)
        ws = wb["VOC Collection"]
        # title row preserved
        assert ws["A1"].value == "Paytm UPI VOC collection template"
        # header still at row 4, extra columns appended after the 12 template cols
        assert ws.cell(row=4, column=1).value == "VOC ID"
        assert ws.cell(row=4, column=13).value == "Group"
        # sample VOC-001 cleared; first data row is our first respondent's id
        first_id = ws.cell(row=5, column=1).value
        assert first_id in set(respondents["voc_id"])
        # one data row per respondent
        data_rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
        assert len(data_rows) == len(respondents)
