"""
build_summary.py — the 2-page executive summary (docx -> PDF).

Layout (A12):
  Page 1 Problem: headline stat, segment & occasions, Chart 1, Chart 2,
                  switching triggers, sizing from observed leakage, 3 quotes.
  Page 2 Solution: chosen solution (A7), finding->feature table, value prop,
                  why Paytm is positioned (evidence only), pilot, a SEPARATE
                  line for potentially addressable users.

Rules enforced here:
  - Every number is inserted from metrics.json (missing -> "—", never invented).
  - Only VERIFIED features (from upi_circle_verification.md) are stated as
    existing; others are framed as proposals.
  - Quotes come from docs/quote_shortlist.md "Selected" list (real verbatims).
  - Exactly 2 pages is enforced later by validate.py.
"""
from __future__ import annotations

import os
import re
import sys
import json

from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from ingest import load_config  # noqa: E402


def pct(x):
    return f"{round(x * 100)}%" if isinstance(x, (int, float)) else "—"


def num(x):
    return f"{x:,}" if isinstance(x, (int, float)) else "—"


def verified_claims(path: str) -> list[str]:
    """Return claim texts whose Status column == VERIFIED."""
    if not os.path.exists(path):
        return []
    claims = []
    for line in open(path, encoding="utf-8"):
        if not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) >= 3 and cells[2].upper() == "VERIFIED":
            claims.append(cells[1])
    return claims


def selected_quotes(path: str) -> list[str]:
    if not os.path.exists(path):
        return []
    text = open(path, encoding="utf-8").read()
    m = re.search(r"##\s*Selected.*", text, re.S)
    if not m:
        return []
    quotes = []
    for line in m.group(0).splitlines()[1:]:
        s = re.sub(r"^\s*\d+\.\s*", "", line).strip()
        if s:
            quotes.append(s)
    return quotes[:3]


def build_summary(metrics: dict, verification_path: str, quotes_path: str,
                  charts_dir: str, out_docx: str, team: str) -> str:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)

    ps_all = metrics.get("paytm_share", {}).get("all", {})
    ci = metrics.get("concept_interest", {})
    siz = metrics.get("sizing", {})
    meta = metrics.get("_meta", {})

    # ---------- Page 1: Problem ----------
    doc.add_heading(f"{team} — Paytm UPI Growth: Problem & Opportunity", level=0)
    head = (f"Among {num(meta.get('n_voc'))} college UPI users sampled, Paytm "
            f"carries {pct(ps_all.get('by_value'))} of payment value and "
            f"{pct(ps_all.get('by_count'))} by count.")
    doc.add_paragraph(head)

    doc.add_heading("Segment & occasions", level=2)
    doc.add_paragraph(
        "College students who made a UPI payment in the last 30 days, split into "
        "Paytm-primary (G1) and other-primary (G2/G3). Occasions: merchant, P2P, "
        "bills, recharge.")

    for fname, cap in (("chart1_share_by_occasion.png", "Chart 1 — Paytm share by occasion (count vs value)"),
                       ("chart2_barrier_codes.png", "Chart 2 — Unprompted barriers")):
        p = os.path.join(charts_dir, fname)
        if os.path.exists(p):
            doc.add_picture(p, width=Inches(5.5))
            doc.add_paragraph(cap).italic = True

    doc.add_heading("Switching & leakage", level=2)
    doc.add_paragraph(
        f"Reported primary-app switches: {num(metrics.get('switching', {}).get('n_reported_switch'))}. "
        f"Observed value not on Paytm (leakage): {pct(siz.get('observed_leakage_value_share'))}.")

    doc.add_heading("Quotes", level=2)
    quotes = selected_quotes(quotes_path)
    if quotes:
        for q in quotes:
            doc.add_paragraph(q, style="Intense Quote")
    else:
        doc.add_paragraph("[3 quotes pending selection in docs/quote_shortlist.md]")

    doc.add_page_break()

    # ---------- Page 2: Solution ----------
    doc.add_heading("Solution", level=1)
    doc.add_paragraph(
        "[Chosen solution from docs/decision.md — finalised after Sahaj confirms. "
        "This section states only VERIFIED features as existing.]")

    doc.add_heading("Finding → feature", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Finding (from metrics)"
    table.rows[0].cells[1].text = "Feature response"
    for _ in range(3):
        table.add_row()

    doc.add_heading("Why Paytm is well positioned (evidence only)", level=2)
    parent = metrics.get("parent_uses_paytm", {})
    doc.add_paragraph(
        f"Parents using Paytm (respondents reporting Yes): {num(parent.get('Yes'))}. "
        "Verified Paytm capabilities relevant to the solution:")
    vc = verified_claims(verification_path)
    if vc:
        for c in vc:
            doc.add_paragraph(c, style="List Bullet")
    else:
        doc.add_paragraph("None verified yet — solution framed as a proposal.",
                          style="List Bullet")

    doc.add_heading("Pilot", level=2)
    doc.add_paragraph(
        "One campus, 4 weeks. Compare activated vs invited-not-activated and "
        "before vs after. Primary metric rises from baseline among activated "
        "students. A minimum activation rate (set from the pilot's own baseline) "
        "gates continuation. No fixed multiplier promised.")

    doc.add_heading("Potentially addressable users (separate — not a forecast)", level=2)
    doc.add_paragraph(
        f"Concept interest: {pct(ci.get('pct_rated_4_or_5'))} rated 4–5. "
        f"Observed-only upper bound of switchable users: "
        f"{num(siz.get('addressable_users_upper_bound'))}. "
        "This is not an adoption forecast and is excluded from any GMV figure.")

    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc.save(out_docx)
    return out_docx


def try_pdf(docx_path: str) -> str | None:
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        return pdf_path
    except Exception as e:  # Word not available / headless
        print(f"PDF export failed ({e}). B6 fallback: open the .docx in Word and "
              f"'Save as PDF', then run validate.py on it.")
        return None


def main():
    cfg = load_config()
    paths = cfg["paths"]
    metrics_path = os.path.join(ROOT, paths["metrics"])
    if not os.path.exists(metrics_path):
        print(f"NOTE: {metrics_path} not found — run the pipeline after DATA PULL.")
        return
    metrics = json.load(open(metrics_path, encoding="utf-8"))
    team = cfg.get("team_name", "TeamName")
    out_docx = os.path.join(ROOT, paths["summary_out"], f"{team}_ExecSummary.docx")
    build_summary(
        metrics,
        os.path.join(ROOT, paths["verification"]),
        os.path.join(ROOT, "docs", "quote_shortlist.md"),
        os.path.join(ROOT, paths["charts_dir"]),
        out_docx, team)
    print(f"wrote {out_docx}")
    pdf = try_pdf(out_docx)
    if pdf:
        print(f"wrote {pdf}")


if __name__ == "__main__":
    sys.exit(main())
