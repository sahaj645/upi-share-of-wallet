"""
validate.py — submission gate. Exit code 0 = pass, 1 = fail.

Checks:
  - file names match the convention ({team}_VOC.xlsx, {team}_ExecSummary.pdf)
  - PDF is exactly 2 pages
  - workbook has >=50 unique VOCs, >=10 in-depth, all groups (G1/G2/G3) present
  - no emails / phone numbers / UPI IDs anywhere in the deliverables (regex)
  - every number in the PDF text appears in outputs/metrics.json (or is a
    whitelisted structural constant)
  - the word SYNTHETIC appears nowhere in the deliverables
"""
from __future__ import annotations

import os
import re
import sys
import json

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from ingest import load_config  # noqa: E402

EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE = re.compile(r"(?<!\d)(?:\+?91[-\s]?)?[6-9]\d{9}(?!\d)")
UPI_ID = re.compile(r"\b[\w.\-]{2,}@(?:ok\w+|paytm|ybl|axl|upi|hdfc\w*|sbi\w*|ibl)\b", re.I)
MIN_VOC = 50
MIN_INDEPTH = 10
GROUPS = {"G1", "G2", "G3"}


def _num_tokens(text: str) -> set[str]:
    return set(re.findall(r"\d[\d,]*(?:\.\d+)?%?", text))


def _allowed_numbers(metrics: dict, cfg: dict) -> set[str]:
    allowed = set()

    def add(v):
        if isinstance(v, bool):
            return
        if isinstance(v, (int, float)):
            allowed.add(str(v))
            allowed.add(f"{v:,}")
            allowed.add(f"{round(v * 100)}%")
            allowed.add(f"{round(v)}")
        elif isinstance(v, dict):
            for x in v.values():
                add(x)
        elif isinstance(v, list):
            for x in v:
                add(x)

    add(metrics)
    # structural constants that legitimately appear in fixed template text
    for c in [cfg.get("small_ticket_max"), cfg.get("aishe_enrolment"),
              1, 2, 3, 4, 5, 10, 90, 30]:
        allowed.add(str(c))
        if isinstance(c, int):
            allowed.add(f"{c:,}")
    return allowed


def _pdf_text_and_pages(path: str):
    from pypdf import PdfReader
    r = PdfReader(path)
    text = "\n".join((p.extract_text() or "") for p in r.pages)
    return text, len(r.pages)


def _xlsx_text(path: str) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.extend(str(c) for c in row if c is not None)
    wb.close()
    return "\n".join(out)


def run_checks(deliverables_dir: str, metrics: dict, cfg: dict) -> list[tuple[str, bool, str]]:
    team = cfg.get("team_name", "TeamName")
    voc = os.path.join(deliverables_dir, f"{team}_VOC.xlsx")
    pdf = os.path.join(deliverables_dir, f"{team}_ExecSummary.pdf")
    results = []

    def check(name, ok, detail=""):
        results.append((name, bool(ok), detail))

    # naming
    check("VOC workbook named {team}_VOC.xlsx", os.path.exists(voc), voc)
    check("Exec summary named {team}_ExecSummary.pdf", os.path.exists(pdf),
          pdf if os.path.exists(pdf) else "missing — export the .docx to PDF (B6)")

    # workbook checks
    if os.path.exists(voc):
        wb = load_workbook(voc, read_only=True, data_only=True)
        ws = wb["VOC Collection"] if "VOC Collection" in wb.sheetnames else wb.worksheets[0]
        header = [str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(header)}
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        ids = {r[0] for r in rows if r and r[0] not in (None, "")}
        check(f"≥{MIN_VOC} unique VOCs", len(ids) >= MIN_VOC, f"found {len(ids)}")
        gi = idx.get("Group")
        ti = idx.get("Interview type")
        groups_present = {r[gi] for r in rows if gi is not None and r[gi]} if gi is not None else set()
        check("All groups G1/G2/G3 present", GROUPS.issubset(groups_present),
              f"found {sorted(groups_present)}")
        indepth = sum(1 for r in rows if ti is not None and str(r[ti]).strip().lower() == "in-depth")
        check(f"≥{MIN_INDEPTH} in-depth interviews", indepth >= MIN_INDEPTH, f"found {indepth}")
        wb.close()

    # gather deliverable text for PII / SYNTHETIC / number checks
    all_text = ""
    pdf_text = ""
    if os.path.exists(voc):
        all_text += _xlsx_text(voc)
    if os.path.exists(pdf):
        pdf_text, pages = _pdf_text_and_pages(pdf)
        all_text += "\n" + pdf_text
        check("PDF is exactly 2 pages", pages == 2, f"{pages} pages")

    # PII scan
    pii = EMAIL.findall(all_text) + PHONE.findall(all_text) + UPI_ID.findall(all_text)
    check("No emails / phones / UPI IDs in deliverables", not pii,
          f"found: {pii[:5]}" if pii else "clean")

    # SYNTHETIC scan
    check("No 'SYNTHETIC' in deliverables", "SYNTHETIC" not in all_text.upper(),
          "found SYNTHETIC" if "SYNTHETIC" in all_text.upper() else "clean")

    # number traceability (PDF only)
    if pdf_text:
        allowed = _allowed_numbers(metrics, cfg)
        stray = sorted(t for t in _num_tokens(pdf_text) if t not in allowed
                       and t.rstrip("%") not in allowed)
        check("Every number in the PDF traces to metrics.json", not stray,
              f"untraceable: {stray}" if stray else "all traceable")

    return results


def main():
    cfg = load_config()
    metrics_path = os.path.join(ROOT, cfg["paths"]["metrics"])
    metrics = json.load(open(metrics_path, encoding="utf-8")) if os.path.exists(metrics_path) else {}
    deliv = os.path.join(ROOT, cfg["paths"]["workbook_out"])
    results = run_checks(deliv, metrics, cfg)

    failed = 0
    for name, ok, detail in results:
        mark = "PASS" if ok else "FAIL"
        if not ok:
            failed += 1
        print(f"[{mark}] {name}" + (f"  — {detail}" if detail else ""))
    print(f"\n{len(results) - failed}/{len(results)} checks passed.")
    if failed:
        print("Gate FAILED. Fix the items above before tagging v1.0-submission.")
        return 1
    print("Gate PASSED. Tag: git tag v1.0-submission && git push --tags")
    return 0


if __name__ == "__main__":
    sys.exit(main())
