"""
build_workbook.py — assemble deliverables/{team_name}_VOC.xlsx.

Sheets (A12):
  1. VOC Collection  — original template columns (order preserved) + extra cols
  2. Payments        — one row per payment
  3. Summary         — headline metrics + embedded charts
  4. Codebook        — from docs/codebook.yaml
  5. Method          — sample, groups, definitions, honest n per group

If reference/Paytm_UPI_VOC_Template.xlsx exists, its 'VOC Collection' header row
is preserved. Otherwise a standard 12-column VOC layout is used (replace with the
official template's columns when available).

No fabricated rows: data comes only from data/clean + outputs/metrics.json.
"""
from __future__ import annotations

import os
import sys
import json

import yaml
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from ingest import load_config  # noqa: E402

# fallback 12-column VOC layout (used only if the official template is absent)
DEFAULT_VOC_COLS = [
    "VOC ID", "Segment", "Date", "Channel", "Prompt / Question",
    "Verbatim (exact words)", "Theme", "Sentiment", "Pain point",
    "Need / Job-to-be-done", "Opportunity", "Priority",
]
EXTRA_COLS = [
    "Group", "Year", "Hostel/Day", "Home state", "Interview type",
    "Paytm share last10 (#)", "Paytm share last10 (₹)", "Parent-funded",
    "Parent uses Paytm", "Concept score", "Visibility preference",
    "Barrier codes", "Switch history",
]


def _voc_template_cols(template_path: str) -> list[str]:
    if os.path.exists(template_path):
        wb = load_workbook(template_path, read_only=True)
        ws = wb["VOC Collection"] if "VOC Collection" in wb.sheetnames else wb.worksheets[0]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        wb.close()
        header = [h for h in header if h is not None]
        if header:
            return header
    return list(DEFAULT_VOC_COLS)


def _per_respondent_paytm(payments: pd.DataFrame) -> dict:
    """voc_id -> (paytm share by count, paytm share by value) over their last 10."""
    out = {}
    if not len(payments):
        return out
    for vid, sub in payments.groupby("voc_id"):
        is_p = sub["app"].astype(str).str.strip().str.lower().eq("paytm")
        cnt = round(float(is_p.sum()) / len(sub), 3) if len(sub) else None
        val_total = pd.to_numeric(sub["amount"], errors="coerce").fillna(0).sum()
        val_p = pd.to_numeric(sub.loc[is_p, "amount"], errors="coerce").fillna(0).sum()
        val = round(float(val_p) / float(val_total), 3) if val_total else None
        out[vid] = (cnt, val)
    return out


def build_voc_sheet(wb, respondents, payments, template_cols, coding):
    ws = wb.create_sheet("VOC Collection")
    ws.append(list(template_cols) + EXTRA_COLS)
    share = _per_respondent_paytm(payments)
    codes_by_vid = {}
    if coding is not None and len(coding):
        for _, r in coding.iterrows():
            fc = str(r.get("final_code") or "").strip()
            if fc:
                codes_by_vid.setdefault(r["voc_id"], []).append(fc)
    for _, r in respondents.iterrows():
        cnt, val = share.get(r.get("voc_id"), (None, None))
        base = [""] * len(template_cols)
        if template_cols:
            base[0] = r.get("voc_id")  # first col is an ID by convention
        extra = [
            r.get("group"), r.get("year"), r.get("hostel_day"), r.get("home_state"),
            "In-depth" if r.get("is_indepth") in (True, "True") else "Survey",
            cnt, val, r.get("funding"), r.get("parent_uses_paytm"),
            r.get("concept_score"), r.get("visibility_pref"),
            ", ".join(codes_by_vid.get(r.get("voc_id"), [])), r.get("switch_history"),
        ]
        ws.append(base + extra)
    return ws


def build_payments_sheet(wb, payments):
    ws = wb.create_sheet("Payments")
    cols = ["voc_id", "n", "amount", "type", "app", "why", "is_small"]
    ws.append(cols)
    for _, r in payments.iterrows():
        ws.append([r.get(c) for c in cols])
    return ws


def build_summary_sheet(wb, metrics, charts_dir):
    ws = wb.create_sheet("Summary")
    ws.append(["Metric", "Value"])
    meta = metrics.get("_meta", {})
    ws.append(["Respondents", meta.get("n_respondents")])
    ws.append(["VOCs", meta.get("n_voc")])
    ws.append(["In-depth interviews", meta.get("n_indepth")])
    ws.append(["Payments captured", meta.get("n_payments")])
    ps = metrics.get("paytm_share", {}).get("all", {})
    ws.append(["Paytm share by count (all)", ps.get("by_count")])
    ws.append(["Paytm share by value (all)", ps.get("by_value")])
    ci = metrics.get("concept_interest", {})
    ws.append(["Concept interest — pct 4/5 (NOT a forecast)", ci.get("pct_rated_4_or_5")])
    siz = metrics.get("sizing", {})
    ws.append(["Addressable users upper bound (observed only)",
               siz.get("addressable_users_upper_bound")])
    ws.append([])
    ws.append(["All figures above are computed in outputs/metrics.json."])
    # embed charts if present
    row = ws.max_row + 2
    for fname in ("chart1_share_by_occasion.png", "chart2_barrier_codes.png"):
        p = os.path.join(charts_dir, fname)
        if os.path.exists(p):
            img = XLImage(p)
            ws.add_image(img, f"A{row}")
            row += 22
    return ws


def build_codebook_sheet(wb, codebook):
    ws = wb.create_sheet("Codebook")
    ws.append(["id", "label", "definition"])
    for c in (codebook or {}).get("codes", []):
        ws.append([c.get("id"), c.get("label"), c.get("definition")])
    return ws


def build_method_sheet(wb, respondents, metrics):
    ws = wb.create_sheet("Method")
    lines = [
        ["Track", "A — build primary-app preference"],
        ["Segment", "College students who made a UPI payment in last 30 days"],
        ["Groups", "G1 Paytm-primary; G2 other-primary + Paytm in 90d; G3 other-primary, no Paytm in 90d"],
        ["Small ticket", "<= ₹500 (config small_ticket_max)"],
        ["No fabrication", "All rows from real responses; synthetic data only in tests/"],
        ["Numbers", "Every figure traces to outputs/metrics.json"],
    ]
    for k, v in lines:
        ws.append([k, v])
    ws.append([])
    ws.append(["Actual n per group (report honestly):"])
    for g, n in (metrics.get("group_counts", {}) or {}).items():
        ws.append([g, n])
    return ws


def build_workbook(respondents, payments, metrics, codebook, template_path,
                   out_path, charts_dir):
    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet
    template_cols = _voc_template_cols(template_path)
    build_voc_sheet(wb, respondents, payments, template_cols, _load_coding())
    build_payments_sheet(wb, payments)
    build_summary_sheet(wb, metrics, charts_dir)
    build_codebook_sheet(wb, codebook)
    build_method_sheet(wb, respondents, metrics)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def _load_coding():
    cfg = load_config()
    p = os.path.join(ROOT, cfg["paths"]["coding_sheet"])
    return pd.read_csv(p) if os.path.exists(p) else None


def main():
    cfg = load_config()
    paths = cfg["paths"]
    rp = os.path.join(ROOT, paths["clean_respondents"])
    if not os.path.exists(rp):
        print(f"NOTE: {rp} not found — run the pipeline after DATA PULL first.")
        return
    respondents = pd.read_csv(rp)
    pp = os.path.join(ROOT, paths["clean_payments"])
    payments = pd.read_csv(pp) if os.path.exists(pp) else pd.DataFrame(
        columns=["voc_id", "n", "amount", "type", "app", "why", "is_small"])
    metrics_path = os.path.join(ROOT, paths["metrics"])
    metrics = json.load(open(metrics_path, encoding="utf-8")) if os.path.exists(metrics_path) else {}
    codebook = yaml.safe_load(open(os.path.join(ROOT, paths["codebook"]), encoding="utf-8"))
    team = cfg.get("team_name", "TeamName")
    out = os.path.join(ROOT, paths["workbook_out"], f"{team}_VOC.xlsx")
    build_workbook(respondents, payments, metrics, codebook,
                   os.path.join(ROOT, paths["voc_template"]), out,
                   os.path.join(ROOT, paths["charts_dir"]))
    print(f"wrote {out}")


if __name__ == "__main__":
    sys.exit(main())
